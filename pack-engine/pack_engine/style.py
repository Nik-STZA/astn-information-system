"""Workbook styling - the single place colours and fonts are defined.

Sheet contents are neutral by design (the management pack standard: black
header bars, grey detail lines, no brand colour). Only the sheet tabs carry the
client's brand, from the profile's tab colours, defaulting to the STZA standard.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT = "Calibri"
BLACK = "000000"
WHITE = "FFFFFF"
GREY_TEXT = "71717A"
ZEBRA = "F2F2F2"
NEGATIVE_RED = "8B0000"

# STZA Brand Guidelines v1.0 - used for tabs when a client has no brand of its own.
STZA_TAB_COLOURS = {
    "summary": "1A1C1E",    # Brand Dark
    "dashboard": "C5A059",  # Brand Gold
    "entity": "D4C5A9",     # Gold Border
    "detail": "F5F0E8",     # Warm Light
    "control": "8E9196",    # Warm Grey
}

NUM = '#,##0;(#,##0);"-"'
NUM_2DP = '#,##0.00;(#,##0.00);"-"'

thin = Side(style="thin", color=BLACK)
medium = Side(style="medium", color=BLACK)
double = Side(style="double", color=BLACK)


def font(size=9, bold=False, italic=False, color=BLACK):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)


def fill(colour):
    return PatternFill("solid", fgColor=colour)


# Row roles, from the management pack formatting standard.
ROW_STYLES = {
    "title":         dict(font=font(13, bold=True)),
    "subtitle":      dict(font=font(10, bold=True)),
    "header":        dict(font=font(9, bold=True, color=WHITE), fill=fill(BLACK)),
    "section":       dict(font=font(9, bold=True)),
    "detail":        dict(font=font(8, color=GREY_TEXT), fill=fill(ZEBRA)),
    "subtotal":      dict(font=font(9, bold=True), border=Border(top=thin)),
    "section_total": dict(font=font(9, bold=True), border=Border(top=thin, bottom=medium)),
    "grand_total":   dict(font=font(9, bold=True), border=Border(top=thin, bottom=double)),
    "check":         dict(font=font(8, italic=True), border=Border(top=thin, bottom=double)),
    "note":          dict(font=font(8, italic=True, color=GREY_TEXT)),
}


def apply_row(ws, row: int, role: str, first_col: int, last_col: int, skip_cols=(), text_cols=()):
    st = ROW_STYLES[role]
    for c in range(first_col, last_col + 1):
        if c in skip_cols:
            continue
        cell = ws.cell(row, c)
        if "font" in st:
            cell.font = st["font"]
        if "fill" in st:
            cell.fill = st["fill"]
        if "border" in st:
            cell.border = st["border"]
        if role == "header" and c > first_col and c not in text_cols:
            cell.alignment = Alignment(horizontal="right")
