"""Build a management pack.

    python -m pack_engine.build --client stza --period 2026-08 --out ./drafts

Needs FINANCE_API_URL and FINANCE_API_KEY. Writes the workbook and a
controls JSON next to it. Exit code 0 when every control passes, 3 when the
pack was written but a control failed.
"""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl.styles import Alignment, Font

from . import ENGINE_VERSION, controls as ctl, style
from .api import FinanceApi
from .fiscal import FiscalCalendar, date_label, month_end, parse_period
from .ledger import fetch_ledger
from .mapping import BY_KEY, resolve_mapping
from .model import build_model
from .render import BS, PNL, SRC, render_statements

PROFILES = Path(__file__).resolve().parent.parent / "profiles"

TAB_ROLES = {"Contents": "summary", PNL: "summary", BS: "summary",
             "Controls": "control", "Mapping": "control", SRC: "detail"}


def load_profile(client: str) -> dict:
    path = PROFILES / f"{client}.json"
    if not path.exists():
        raise SystemExit(f"No pack profile for {client!r} ({path})")
    return json.loads(path.read_text(encoding="utf-8"))


@dataclass
class BuildResult:
    path: Path
    controls: list
    failed: list


def _controls_sheet(wb, controls):
    ws = wb.create_sheet("Controls")
    ws["A1"], ws["A2"] = "Controls", "A failed control puts FAILED CONTROLS in the file name."
    ws["A1"].font = style.ROW_STYLES["title"]["font"]
    ws["A2"].font = style.ROW_STYLES["note"]["font"]
    head = ["Control", "Status", "Expected", "Actual", "Difference", "Detail"]
    for i, h in enumerate(head, 1):
        ws.cell(4, i, h)
    style.apply_row(ws, 4, "header", 1, len(head), text_cols=(2, 6))
    for r, c in enumerate(controls, start=5):
        vals = [c.label, c.status.upper(), c.expected, c.actual, c.difference, c.detail]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(r, i, v)
            cell.font = style.font(9, bold=(i == 2 and c.status != "pass"),
                                   color=style.NEGATIVE_RED if (i == 2 and c.status == "fail") else style.BLACK)
            if i in (3, 4, 5):
                cell.number_format = style.NUM_2DP
    for letter, width in zip("ABCDEF", (52, 9, 13, 13, 12, 90)):
        ws.column_dimensions[letter].width = width


def _mapping_sheet(wb, model, mapping):
    ws = wb.create_sheet("Mapping")
    ws["A1"] = "Account mapping"
    ws["A2"] = "Proposed by rule from Xero's account types and reporting codes. Not yet approved."
    ws["A1"].font = style.ROW_STYLES["title"]["font"]
    ws["A2"].font = style.ROW_STYLES["note"]["font"]
    head = ["Code", "Account", "Xero class", "Xero type", "Reporting code", "Statement", "Category", "Why"]
    for i, h in enumerate(head, 1):
        ws.cell(4, i, h)
    style.apply_row(ws, 4, "header", 1, len(head), text_cols=range(1, 9))
    r = 5
    for block in [*model.pnl, *model.bs]:
        for line in block.lines:
            a, m = line.account, mapping[line.account.id]
            cat = BY_KEY[m.category]
            for i, v in enumerate([a.code, a.name, a.klass, a.type, a.reporting_code,
                                   "Profit and loss" if cat.statement == "pnl" else "Balance sheet",
                                   cat.label, m.reason], 1):
                ws.cell(r, i, v).font = style.font(9)
            r += 1
    for g in model.gaps:
        a = g.account
        for i, v in enumerate([a.code, a.name, a.klass, a.type, a.reporting_code, "", "UNMAPPED", g.reason], 1):
            ws.cell(r, i, v).font = style.font(9, bold=True, color=style.NEGATIVE_RED)
        r += 1
    for letter, width in zip("ABCDEFGH", (8, 38, 11, 12, 18, 15, 44, 80)):
        ws.column_dimensions[letter].width = width


def _contents_sheet(wb, profile, period, fy, controls, stamp):
    ws = wb["Contents"]
    y, m = parse_period(period)
    fails = ctl.failed(controls)
    warns = [c for c in controls if c.status == "warn"]
    passed = len(controls) - len(fails) - len(warns)
    summary = f"Controls: {passed} passed"
    if warns:
        summary += f", {len(warns)} warning{'s' if len(warns) != 1 else ''}"
    if fails:
        summary += f", {len(fails)} FAILED - see Controls"
    lines = [
        (profile["legal_name"], "title"),
        (f"Management pack - {month_end(y, m).strftime('%B %Y')} ({fy})", "subtitle"),
        ("", None),
        ("DRAFT for review. A management presentation, not statutory accounts.", "note"),
        (f"Framework {profile['framework']}; currency {profile['currency']}; financial year ends "
         f"{month_end(2001, profile['year_end_month']).strftime('%d %B').lstrip('0')}.", "note"),
        (f"Built {stamp} by pack engine {ENGINE_VERSION} from Xero via stza-finance-api.", "note"),
        ("Account mapping proposed by rule, not yet approved.", "note"),
        ("", None),
        (summary, "subtitle"),
        ("", None),
        ("Contents", "subtitle"),
    ]
    for r, (text, role) in enumerate(lines, 1):
        ws.cell(r, 1, text)
        if role:
            ws.cell(r, 1).font = style.ROW_STYLES[role]["font"]
    r = len(lines) + 1
    for name in wb.sheetnames:
        if name == "Contents":
            continue
        cell = ws.cell(r, 1, name)
        cell.hyperlink = f"#'{name}'!A1"
        cell.font = Font(name=style.FONT, size=10, underline="single", color=style.BLACK)
        r += 1
    ws.column_dimensions["A"].width = 100
    ws.sheet_view.showGridLines = False


def build(client: str, period: str, out_dir: Path, api: FinanceApi | None = None,
          now: datetime | None = None) -> BuildResult:
    profile = load_profile(client)
    api = api or FinanceApi()
    cal = FiscalCalendar(profile["year_end_month"])
    fy, prior_fy = cal.label(period), cal.prior_label(period)
    y, m = parse_period(period)
    stamp_dt = now or datetime.now()
    stamp = stamp_dt.strftime("%d %b %Y %H:%M").lstrip("0")

    ledger = fetch_ledger(api, client, profile["entity"], cal, period)
    mapping = resolve_mapping(ledger.active_accounts())
    model = build_model(ledger, mapping)
    rendered = render_statements(model, ledger, client_label=profile["legal_name"], fy=fy,
                                 prior_fy=prior_fy, period_month=month_end(y, m))

    controls = [*ctl.mapping_controls(ledger, model), *ctl.ledger_controls(ledger),
                *ctl.balance_controls(model),
                *ctl.xero_tie_controls(api, client, profile["entity"], cal, period, model),
                *ctl.formula_controls(rendered)]

    wb = rendered.workbook
    _controls_sheet(wb, controls)
    _mapping_sheet(wb, model, mapping)
    wb.move_sheet(SRC, offset=len(wb.sheetnames))
    _contents_sheet(wb, profile, period, fy, controls, stamp)

    tab_colours = (profile.get("branding") or {}).get("tab_colours") or style.STZA_TAB_COLOURS
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = tab_colours[TAB_ROLES.get(ws.title, "detail")]

    fails = ctl.failed(controls)
    tag = month_end(y, m).strftime("%b%Y")
    name = (f"{profile['label']} - Management Pack - {tag} - draft {stamp_dt.strftime('%Y-%m-%d %H%M')}"
            + (" - FAILED CONTROLS" if fails else "") + ".xlsx")
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / name
    wb.save(path)
    (out_dir / (path.stem + " - controls.json")).write_text(
        json.dumps({"client": client, "period": period, "engine_version": ENGINE_VERSION,
                    "controls": [c.as_dict() for c in controls]}, indent=2), encoding="utf-8")
    return BuildResult(path, controls, fails)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Build a management pack")
    ap.add_argument("--client", required=True)
    ap.add_argument("--period", required=True, help="YYYY-MM")
    ap.add_argument("--out", required=True)
    args = ap.parse_args(argv)
    result = build(args.client, args.period, Path(args.out))
    for c in result.controls:
        print(f"[{c.status.upper():4}] {c.label}" + (f" - {c.detail}" if c.status != "pass" and c.detail else ""))
    print(f"wrote {result.path}")
    return 3 if result.failed else 0


if __name__ == "__main__":
    sys.exit(main())
