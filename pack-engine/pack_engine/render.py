"""Workbook renderer.

Rules:
  - The only typed numbers are on "Source - Trial balance" (Xero trial balance
    figures, debit minus credit). Every figure on the statements is a formula
    that points at it, and every subtotal and total is a formula.
  - No hardcoded row numbers anywhere: rows are assigned while writing, and the
    rows other sheets need are recorded as anchors and Excel defined names.
  - Months after the reporting month are left blank.
Column layout on both statements: A label | B spacer | C/D/E headline columns |
F spacer | G-R the twelve months of the financial year.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.defined_name import DefinedName

from . import style
from .fiscal import date_label, month_label
from .mapping import BY_KEY
from .model import CREDIT_SECTIONS, Model

SRC = "Source - Trial balance"
PNL = "Profit and loss"
BS = "Balance sheet"
MONTH_COL0 = 7          # G
HEAD = 4                # header row on the statements


@dataclass
class Rendered:
    workbook: Workbook
    anchors: dict[str, int] = field(default_factory=dict)       # "pnl.net_profit" -> row number
    formula_cells: list[tuple[str, str]] = field(default_factory=list)   # (sheet, coord) that must hold formulas


def _q(sheet: str) -> str:
    return f"'{sheet}'"


def _month_cols(model: Model) -> dict[date, str]:
    return {d: col(MONTH_COL0 + i) for i, d in enumerate(model.fy_months)}


class _Writer:
    def __init__(self, out: Rendered, ws, title: str, subtitle: str, headers: list[str]):
        self.out, self.ws, self.r = out, ws, HEAD + 1
        ws["A1"], ws["A2"] = title, subtitle
        ws["A1"].font, ws["A2"].font = style.ROW_STYLES["title"]["font"], style.ROW_STYLES["subtitle"]["font"]
        for i, h in enumerate(headers, start=1):
            if h is not None:
                ws.cell(HEAD, i, h)
        style.apply_row(ws, HEAD, "header", 1, len(headers), skip_cols=(2, 6))
        ws.column_dimensions["A"].width = 46
        for c in ("B", "F"):
            ws.column_dimensions[c].width = 1.6
        for i in range(3, len(headers) + 1):
            if i not in (2, 6):
                ws.column_dimensions[col(i)].width = 12
        ws.freeze_panes = "C5"

    def row(self, label: str, role: str, cells: dict[str, object], indent: int = 0, formula_required=False):
        r = self.r
        self.ws.cell(r, 1, ("  " * indent) + label)
        for c, v in cells.items():
            cell = self.ws[f"{c}{r}"]
            cell.value = v
            cell.number_format = style.NUM
            if formula_required and v is not None:
                self.out.formula_cells.append((self.ws.title, f"{c}{r}"))
        style.apply_row(self.ws, r, role, 1, MONTH_COL0 + 11, skip_cols=(2, 6))
        self.r += 1
        return r

    def blank(self):
        self.r += 1


def _source_sheet(out: Rendered, model: Model) -> dict[str, int]:
    """Write the trial balance tab; returns account id -> row."""
    ws = out.workbook.create_sheet(SRC)
    ws["A1"] = "Source - Trial balance"
    ws["A1"].font = style.ROW_STYLES["title"]["font"]
    ws["A2"] = ("Typed values from Xero trial balances via stza-finance-api, debit minus credit. "
                "Movement is the month; closing is the balance sheet balance, or the year to date "
                "for profit and loss accounts. Every statement figure is a formula over this tab.")
    ws["A2"].font = style.ROW_STYLES["note"]["font"]
    months = model.fy_months
    head = ["Code", "Account", "Xero class", "Xero type", "Reporting code", "Category", "Xero account ID"]
    head += [f"Movement {month_label(d)}" for d in months]
    head += [f"Closing {month_label(d)}" for d in months]
    head += [f"{date_label(model.prior_year_end)} (closing or YTD)"]
    for i, h in enumerate(head, start=1):
        ws.cell(4, i, h)
    style.apply_row(ws, 4, "header", 1, len(head), text_cols=range(1, 8))
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["G"].width = 38
    rows: dict[str, int] = {}
    r = 5
    blocks = [*model.pnl, *model.bs]
    for b in blocks:
        for line in b.lines:
            a = line.account
            ws.cell(r, 1, a.code)
            ws.cell(r, 2, a.name)
            ws.cell(r, 3, a.klass)
            ws.cell(r, 4, a.type)
            ws.cell(r, 5, a.reporting_code)
            ws.cell(r, 6, BY_KEY[b.category].label)
            ws.cell(r, 7, a.id)
            rows[a.id] = r
            r += 1
    return rows


def _src_ref(model: Model, src_rows, account_id: str, month: date | None, kind: str) -> str:
    """kind: 'movement' | 'closing' | 'prior'."""
    n = len(model.fy_months)
    if kind == "prior":
        c = 8 + 2 * n
    else:
        i = model.fy_months.index(month)
        c = 8 + i if kind == "movement" else 8 + n + i
    return f"{_q(SRC)}!{col(c)}{src_rows[account_id]}"


def _fill_source_values(out: Rendered, model: Model, ledger, src_rows):
    ws = out.workbook[SRC]
    n = len(model.fy_months)
    for acct_id, r in src_rows.items():
        for i, d in enumerate(model.fy_months):
            if d in model.months:
                ws.cell(r, 8 + i, ledger.movement(acct_id, d)).number_format = style.NUM_2DP
                ws.cell(r, 8 + n + i, ledger.closing(acct_id, d)).number_format = style.NUM_2DP
        ws.cell(r, 8 + 2 * n, ledger.closing(acct_id, model.prior_year_end)).number_format = style.NUM_2DP


def _pnl_sheet(out: Rendered, model: Model, src_rows, client_label: str, fy: str, prior_fy: str, period_month: date):
    ws = out.workbook.create_sheet(PNL)
    mcols = _month_cols(model)
    pcol = mcols[period_month]
    headers = ["", None, month_label(period_month), f"YTD {fy}", prior_fy, None] + [month_label(d) for d in model.fy_months]
    w = _Writer(out, ws, client_label, "Profit and loss (GBP) - management presentation, not statutory accounts", headers)

    def headline(r):
        return {"C": f"={pcol}{r}", "D": f"=SUM(G{r}:{pcol}{r})"}

    subtotal_rows: dict[str, int] = {}

    def block(b):
        w.row(b.label, "section", {})
        first = w.r
        for line in b.lines:
            r = w.r
            cells = {mcols[d]: f"=-{_src_ref(model, src_rows, line.account.id, d, 'movement')}" for d in model.months}
            cells.update(headline(r))
            cells["E"] = f"=-{_src_ref(model, src_rows, line.account.id, None, 'prior')}"
            w.row(line.account.label, "detail", cells, indent=1)
        last = w.r - 1
        cols = ["C", "D", "E", *(mcols[d] for d in model.months)]
        r = w.row(f"Total {b.label[0].lower()}{b.label[1:]}", "subtotal",
                  {c: f"=SUM({c}{first}:{c}{last})" for c in cols}, formula_required=True)
        subtotal_rows[b.category] = r
        w.blank()

    def total(label, keys_or_rows: list, role, key):
        refs = [subtotal_rows[k] if isinstance(k, str) else k for k in keys_or_rows
                if not isinstance(k, str) or k in subtotal_rows]
        cols = ["C", "D", "E", *(mcols[d] for d in model.months)]
        cells = {c: ("=" + "+".join(f"{c}{x}" for x in refs)) if refs else "=0" for c in cols}
        r = w.row(label, role, cells, formula_required=True)
        out.anchors[f"pnl.{key}"] = r
        w.blank()
        return r

    by_section = {}
    for b in model.pnl:
        by_section.setdefault(BY_KEY[b.category].section, []).append(b)

    for b in by_section.get("turnover", []) + by_section.get("cost_of_sales", []):
        block(b)
    gp = total("Gross profit", ["turnover", "cost_of_sales"], "section_total", "gross_profit")
    overheads = by_section.get("overheads", [])
    for b in overheads:
        block(b)
    toh = total("Total overheads", [b.category for b in overheads], "subtotal", "total_overheads")
    op = total("Operating profit", [gp, toh], "section_total", "operating_profit")
    finance = by_section.get("finance", [])
    for b in finance:
        block(b)
    pbt = total("Profit before taxation", [op, *[b.category for b in finance]], "section_total", "profit_before_tax")
    for b in by_section.get("taxation", []):
        block(b)
    total("Profit for the financial year", [pbt, "taxation"], "grand_total", "net_profit")
    _negatives_red(ws, w.r)
    return mcols


def _bs_sheet(out: Rendered, model: Model, src_rows, client_label: str, period_month: date, mcols):
    ws = out.workbook.create_sheet(BS)
    pcol = mcols[period_month]
    idx = model.fy_months.index(period_month)
    prior_col = mcols[model.fy_months[idx - 1]] if idx > 0 else "E"
    headers = ["", None, date_label(period_month),
               date_label(model.fy_months[idx - 1]) if idx > 0 else date_label(model.prior_year_end),
               date_label(model.prior_year_end), None] + [date_label(d) for d in model.fy_months]
    w = _Writer(out, ws, client_label, "Balance sheet (GBP) - management presentation, not statutory accounts", headers)
    cols = ["C", "D", "E", *(mcols[d] for d in model.months)]
    np_row = out.anchors["pnl.net_profit"]
    rows: dict[str, int] = {}

    def headline(r):
        return {"C": f"={pcol}{r}", "D": f"={prior_col}{r}"}

    def block(b):
        sign = "-" if BY_KEY[b.category].section in CREDIT_SECTIONS else ""
        w.row(b.label, "section", {})
        first = w.r
        for line in b.lines:
            r = w.r
            cells = {mcols[d]: f"={sign}{_src_ref(model, src_rows, line.account.id, d, 'closing')}" for d in model.months}
            cells.update(headline(r))
            cells["E"] = f"={sign}{_src_ref(model, src_rows, line.account.id, None, 'prior')}"
            w.row(line.account.label, "detail", cells, indent=1)
        last = w.r - 1
        r = w.row(f"Total {b.label[0].lower()}{b.label[1:]}", "subtotal",
                  {c: f"=SUM({c}{first}:{c}{last})" for c in cols}, formula_required=True)
        rows[b.category] = r
        w.blank()

    def total(label, parts: list[tuple[int, str]], role, key):
        cells = {}
        for c in cols:
            expr = "".join(f"{op}{c}{r}" for r, op in parts)
            cells[c] = "=" + (expr.lstrip("+") if expr else "0")
        r = w.row(label, role, cells, formula_required=True)
        out.anchors[f"bs.{key}"] = r
        w.blank()
        return r

    by_section = {}
    for b in model.bs:
        by_section.setdefault(BY_KEY[b.category].section, []).append(b)
    plus = lambda keys: [(rows[k], "+") for k in keys if k in rows]

    for b in by_section.get("fixed_assets", []):
        block(b)
    fa = total("Total fixed assets", plus(["intangible_assets", "tangible_assets", "investments"]), "subtotal", "fixed_assets")
    for b in by_section.get("current_assets", []):
        block(b)
    ca = total("Total current assets", plus(["stock", "debtors", "cash"]), "subtotal", "current_assets")
    for b in by_section.get("creditors_lt1y", []):
        block(b)
    cl = [(rows["creditors_lt1y"], "-")] if "creditors_lt1y" in rows else []
    nca = total("Net current assets (liabilities)", [(ca, "+"), *cl], "section_total", "net_current_assets")
    talcl = total("Total assets less current liabilities", [(fa, "+"), (nca, "+")], "section_total", "talcl")
    for b in by_section.get("creditors_gt1y", []) + by_section.get("provisions", []):
        block(b)
    na = total("Net assets", [(talcl, "+"), *[(rows[k], "-") for k in ("creditors_gt1y", "provisions") if k in rows]],
               "grand_total", "net_assets")

    w.row("Capital and reserves", "section", {})
    w.blank()
    for b in by_section.get("equity", []):
        block(b)
    cye_cells = {mcols[d]: f"=SUM({_q(PNL)}!$G${np_row}:{mcols[d]}{np_row})" for d in model.months}
    cye_cells.update({"C": f"={pcol}{w.r}", "D": f"={prior_col}{w.r}", "E": f"={_q(PNL)}!$E${np_row}"})
    cye = w.row("Current year earnings", "subtotal", cye_cells, formula_required=True)
    w.blank()
    eq = total("Total equity", [*plus(["share_capital", "retained_earnings", "other_reserves"]), (cye, "+")],
               "grand_total", "total_equity")
    chk = total("Check: net assets less total equity (should be nil)", [(na, "+"), (eq, "-")], "check", "check")
    _negatives_red(ws, w.r)


def _negatives_red(ws, last_row):
    ws.conditional_formatting.add(
        f"C5:R{last_row}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color=style.NEGATIVE_RED)))


def _define_names(out: Rendered, model: Model, period_month: date, mcols):
    wb = out.workbook
    for key, r in out.anchors.items():
        sheet = PNL if key.startswith("pnl.") else BS
        base = key.split(".", 1)[1].upper()
        prefix = "PL" if sheet == PNL else "BS"
        for suffix, c in (("MONTH", "C"), ("YTD" if sheet == PNL else "PRIOR_MONTH", "D"), ("PRIOR", "E")):
            name = f"{prefix}_{base}_{suffix}"
            wb.defined_names[name] = DefinedName(name, attr_text=f"{_q(sheet)}!${c}${r}")


def render_statements(model: Model, ledger, *, client_label: str, fy: str, prior_fy: str, period_month: date) -> Rendered:
    wb = Workbook()
    wb.remove(wb.active)
    out = Rendered(workbook=wb)
    wb.create_sheet("Contents")
    src_rows = _source_sheet(out, model)
    mcols = _pnl_sheet(out, model, src_rows, client_label, fy, prior_fy, period_month)
    _bs_sheet(out, model, src_rows, client_label, period_month, mcols)
    _fill_source_values(out, model, ledger, src_rows)
    # Order: Contents, statements, then controls/mapping (added later), source last.
    wb.move_sheet(SRC, offset=len(wb.sheetnames))
    _define_names(out, model, period_month, mcols)
    return out
