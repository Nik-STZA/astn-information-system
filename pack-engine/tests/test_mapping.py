"""Saved mappings and client categories from finance-api (migration 017)."""

import json
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

import fake_api
from fake_api import FakeApi
from pack_engine import build as build_mod
from pack_engine.api import FinanceApi
from pack_engine.ledger import account_from_xero
from pack_engine.mapping import CATEGORIES, categories_from_api, propose, resolve_mapping

NOW = datetime(2026, 6, 10, 9, 30)
CASES = json.loads((Path(__file__).parent / "mapping_rule_cases.json").read_text(encoding="utf-8"))["cases"]


@pytest.mark.parametrize("case", CASES, ids=[c["note"] for c in CASES])
def test_rules_match_the_cases_shared_with_finance_api(case):
    a = account_from_xero({"AccountID": "x", **case["account"]})
    assert resolve_mapping([a])["x"].category == case["expected"]


def standard_rows(**changes):
    rows = [{"key": c.key, "label": c.label, "statement": c.statement, "section": c.section,
             "sortOrder": c.order, "cashFlow": c.cash_flow, "active": True} for c in CATEGORIES]
    for key, patch in changes.items():
        rows = [{**r, **patch} if r["key"] == key else r for r in rows]
    return rows


def saved(status="approved", categories=None, overrides=None):
    """Every fake account saved with its rule category, as the portal would after 'approve all'."""
    mappings = []
    for a in fake_api.ACCOUNTS:
        acc = account_from_xero(a)
        m = {"accountId": acc.id, "code": acc.code, "name": acc.name, "class": acc.klass,
             "category": propose(acc).category, "cashFlow": None, "status": status, "source": "rule",
             "reason": "rule", "approvedBy": "nik@stza.io" if status == "approved" else None,
             "approvedAt": "2026-09-14T17:00:00Z" if status == "approved" else None}
        mappings.append({**m, **((overrides or {}).get(acc.id, {}))})
    return {"categories": categories if categories is not None else standard_rows(), "mappings": mappings}


@pytest.fixture
def profile(monkeypatch):
    p = {"client": "demo", "label": "DEMO", "legal_name": "Demo Ltd", "entity": "demo",
         "year_end_month": 3, "first_year_end": "2026-03-31", "framework": "FRS 102 Section 1A",
         "currency": "GBP", "branding": {"tab_colours": None}}
    monkeypatch.setattr(build_mod, "load_profile", lambda client: p)
    return p


def control(res, key):
    return next(c for c in res.controls if c.key == key)


def rows_labelled(ws):
    return {str(ws.cell(r, 1).value).strip(): r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value}


def test_a_fully_approved_mapping_passes_the_approval_control(tmp_path, profile):
    res = build_mod.build("demo", "2026-05", tmp_path, api=FakeApi(saved_mapping=saved()), now=NOW)
    assert res.failed == []
    assert control(res, "mapping.approved").status == "pass"
    wb = openpyxl.load_workbook(res.path)
    assert "approved in the portal" in " ".join(str(wb["Contents"].cell(r, 1).value) for r in range(1, 12))
    statuses = {wb["Mapping"].cell(r, 9).value for r in range(5, wb["Mapping"].max_row + 1)} - {None}
    assert statuses == {"Approved"}


def test_saved_but_unapproved_and_rule_only_accounts_warn(tmp_path, profile):
    data = saved(status="proposed")
    data["mappings"] = [m for m in data["mappings"] if m["accountId"] != "a-rent"]    # nothing saved for rent
    res = build_mod.build("demo", "2026-05", tmp_path, api=FakeApi(saved_mapping=data), now=NOW)
    c = control(res, "mapping.approved")
    assert c.status == "warn" and "saved but not approved" in c.detail and "1 mapped by rule only" in c.detail
    assert res.failed == []


def test_an_older_finance_api_falls_back_to_rules_with_a_warning(tmp_path, profile):
    res = build_mod.build("demo", "2026-05", tmp_path, api=FakeApi(saved_mapping=None), now=NOW)
    c = control(res, "mapping.approved")
    assert c.status == "warn" and "no saved mappings" in c.detail
    assert res.failed == []


def test_the_saved_category_wins_over_the_rule(tmp_path, profile):
    data = saved(overrides={"a-rent": {"category": "administrative"}})
    res = build_mod.build("demo", "2026-05", tmp_path, api=FakeApi(saved_mapping=data), now=NOW)
    ws = openpyxl.load_workbook(res.path)["Profit and loss"]
    rows = rows_labelled(ws)
    assert "Establishment costs" not in rows
    assert rows["469 - Rent"] > rows["Administrative expenses"]
    assert res.failed == []


def test_a_category_switched_off_for_the_client_leaves_a_gap_that_fails(tmp_path, profile):
    data = saved(categories=standard_rows(establishment={"active": False}))
    res = build_mod.build("demo", "2026-05", tmp_path, api=FakeApi(saved_mapping=data), now=NOW)
    gap = next(c for c in res.failed if c.key == "mapping.coverage")
    assert "469 - Rent" in gap.detail


def test_a_client_only_category_is_drawn_in_its_section_and_the_statements_still_tie(tmp_path, profile):
    rows = standard_rows() + [{"key": "pension_costs", "label": "Pension costs", "statement": "pnl",
                               "section": "overheads", "sortOrder": 35, "cashFlow": None, "active": True}]
    data = saved(categories=rows, overrides={"a-pen": {"category": "pension_costs"}})
    res = build_mod.build("demo", "2026-05", tmp_path, api=FakeApi(saved_mapping=data), now=NOW)
    assert res.failed == []
    ws = openpyxl.load_workbook(res.path)["Profit and loss"]
    labels = rows_labelled(ws)
    assert labels["Staff costs"] < labels["Pension costs"] < labels["Establishment costs"] < labels["Total overheads"]
    toh = ws.cell(labels["Total overheads"], 6).value
    assert f"F{labels['Pension costs']}" in toh


def test_a_cash_flow_override_moves_one_account(tmp_path, profile):
    data = saved(overrides={"a-cred": {"cashFlow": "loans"}})
    res = build_mod.build("demo", "2026-05", tmp_path, api=FakeApi(saved_mapping=data), now=NOW)
    assert res.failed == []
    wb = openpyxl.load_workbook(res.path)
    cf = rows_labelled(wb["Cash flow"])
    bs_row = rows_labelled(wb["Balance sheet"])["800 - Accounts Payable"]
    assert f"'Balance sheet'!G{bs_row}" in wb["Cash flow"].cell(cf["Loans"], 7).value
    assert "Increase/(decrease) in creditors" not in cf     # nothing else is working capital credit


def test_categories_the_pack_cannot_draw_are_refused():
    with pytest.raises(ValueError, match="cannot draw"):
        categories_from_api(standard_rows(turnover={"section": "fixed_assets"}))


def test_account_mapping_is_none_when_finance_api_has_no_such_route():
    class Resp:
        status_code, text = 404, "Cannot GET"

    class Session:
        def get(self, *a, **k):
            return Resp()

    api = FinanceApi("https://api.example", "k", session=Session(), sleep=lambda s: None)
    assert api.account_mapping("demo", "demo") is None
