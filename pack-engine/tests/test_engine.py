from datetime import date, datetime

import openpyxl
import pytest

from fake_api import FakeApi, acct
from pack_engine import build as build_mod
from pack_engine.fiscal import FiscalCalendar
from pack_engine.ledger import account_from_xero, parse_trial_balance
from pack_engine.mapping import propose, resolve_mapping

NOW = datetime(2026, 6, 10, 9, 30)


# --- financial year -------------------------------------------------------------

def test_march_year_end():
    cal = FiscalCalendar(3)
    months = cal.fy_months("2026-08")
    assert months[0] == date(2026, 4, 30) and months[-1] == date(2027, 3, 31)
    assert cal.months_to_date("2026-08")[-1] == date(2026, 8, 31)
    assert cal.months_to_date("2027-02") == months[:11]
    assert cal.prior_year_end("2026-08") == date(2026, 3, 31)
    assert (cal.label("2026-08"), cal.prior_label("2026-08")) == ("FY27", "FY26")
    assert cal.label("2027-03") == "FY27" and cal.label("2027-04") == "FY28"


def test_december_year_end():
    cal = FiscalCalendar(12)
    assert cal.fy_months("2026-07")[0] == date(2026, 1, 31)
    assert cal.prior_year_end("2026-07") == date(2025, 12, 31)
    assert cal.label("2026-07") == "FY26"


# --- trial balance parsing -----------------------------------------------------

def test_trial_balance_is_movement_and_ytd_by_account_id():
    tb = parse_trial_balance(FakeApi().trial_balance("c", "e", date(2026, 5, 31)), date(2026, 5, 31))
    assert tb.movement["a-rev"] == -900.0          # the month
    assert tb.ytd["a-rev"] == -1700.0              # April + May
    assert tb.ytd["a-bank"] == 1390.0              # closing balance
    assert tb.total_debit == tb.total_credit and tb.total_ytd_debit == tb.total_ytd_credit
    assert tb.unidentified == []


# --- mapping ----------------------------------------------------------------------

@pytest.mark.parametrize("a,expected", [
    (acct("x", "600", "Natwest", "ASSET", "BANK", "ASS"), "cash"),                  # type beats a bare code
    (acct("x", "481", "Pension", "EXPENSE", "OVERHEADS", "EXP"), "staff_costs"),   # UK payroll range
    (acct("x", "201", "Fees", "REVENUE", "REVENUE", "REV"), "turnover"),           # bare REV -> type
    (acct("x", "270", "Interest", "REVENUE", "REVENUE", "REV.INV.INT"), "other_income"),
    (acct("x", "498", "FX", "EXPENSE", "EXPENSE", "EXP.ADM.FOR.UGL"), "fx"),
    (acct("x", "441", "Legal", "EXPENSE", "OVERHEADS", "EXP.ADM.FEE"), "professional_fees"),
    (acct("x", "839", "Dividends payable", "LIABILITY", "CURRLIAB", "LIA.CUR"), "creditors_lt1y"),
    (acct("x", "772", "Investment in X", "ASSET", "NONCURRENT", "ASS.NCA"), "investments"),
])
def test_mapping_rules(a, expected):
    assert propose(account_from_xero(a)).category == expected


def test_an_account_with_no_rule_is_a_gap_not_a_guess():
    m = resolve_mapping([account_from_xero(acct("x", "999", "Odd", "ASSET", "MYSTERY", ""))])
    assert m["x"].is_gap


def test_a_rule_that_crosses_statements_is_refused():
    odd = account_from_xero(acct("x", "999", "Odd", "LIABILITY", "CURRLIAB", "EXP.ADM"))
    assert resolve_mapping([odd])["x"].is_gap


# --- whole build -----------------------------------------------------------------

@pytest.fixture
def profile(monkeypatch):
    p = {"client": "demo", "label": "DEMO", "legal_name": "Demo Ltd", "entity": "demo",
         "year_end_month": 3, "framework": "FRS 102 Section 1A", "currency": "GBP",
         "branding": {"tab_colours": None}}
    monkeypatch.setattr(build_mod, "load_profile", lambda client: p)
    return p


def test_a_clean_build_passes_every_control_and_ties_to_xero(tmp_path, profile):
    res = build_mod.build("demo", "2026-05", tmp_path, api=FakeApi(), now=NOW)
    assert [c.label for c in res.failed] == []
    assert res.path.name == "DEMO - Management Pack - May2026 - draft 2026-06-10 0930.xlsx"
    wb = openpyxl.load_workbook(res.path)
    assert wb.sheetnames == ["Contents", "Profit and loss", "Balance sheet", "Cash flow", "Controls",
                             "Mapping", "Source - Trial balance"]
    assert (tmp_path / (res.path.stem + " - controls.json")).exists()


def test_statements_hold_no_typed_numbers(tmp_path, profile):
    res = build_mod.build("demo", "2026-05", tmp_path, api=FakeApi(), now=NOW)
    wb = openpyxl.load_workbook(res.path)
    for name in ("Profit and loss", "Balance sheet", "Cash flow"):
        typed = [c.coordinate for row in wb[name].iter_rows(min_row=5) for c in row
                 if isinstance(c.value, (int, float))]
        assert typed == [], f"{name} has typed numbers at {typed[:5]}"


def test_statement_lines_point_at_the_source_tab(tmp_path, profile):
    res = build_mod.build("demo", "2026-05", tmp_path, api=FakeApi(), now=NOW)
    ws = openpyxl.load_workbook(res.path)["Profit and loss"]
    sales = next(r for r in range(5, ws.max_row + 1) if str(ws.cell(r, 1).value).strip() == "200 - Sales")
    assert ws.cell(sales, 7).value.startswith("=-'Source - Trial balance'!")     # April, credit shown positive
    assert ws.cell(sales, 4).value == f"=SUM(G{sales}:H{sales})"                  # YTD to May
    assert ws.cell(sales, 9).value is None                                        # June not reported yet


def test_a_mismatch_with_xero_fails_and_says_so_in_the_file_name(tmp_path, profile):
    res = build_mod.build("demo", "2026-05", tmp_path, api=FakeApi(xero_net_profit_offset=1.0), now=NOW)
    assert "FAILED CONTROLS" in res.path.name
    assert {c.key for c in res.failed} >= {"xero.pnl.month", "xero.pnl.ytd", "xero.pnl.prior"}


def test_an_unmapped_account_with_activity_fails(tmp_path, profile, monkeypatch):
    import fake_api
    odd = acct("a-odd", "999", "Mystery", "ASSET", "MYSTERY", "")
    api = FakeApi(extra_account=odd)
    months = {**fake_api.MONTHS[date(2026, 5, 31)], "a-odd": 25.0, "a-bank": 495.0}
    monkeypatch.setitem(fake_api.MONTHS, date(2026, 5, 31), months)
    res = build_mod.build("demo", "2026-05", tmp_path, api=api, now=NOW)
    gap = next(c for c in res.failed if c.key == "mapping.coverage")
    assert "999 - Mystery" in gap.detail


def test_tabs_take_the_brand_colours_and_default_to_the_stza_standard(tmp_path, profile):
    res = build_mod.build("demo", "2026-05", tmp_path, api=FakeApi(), now=NOW)
    wb = openpyxl.load_workbook(res.path)
    assert wb["Profit and loss"].sheet_properties.tabColor.rgb.endswith("1A1C1E")
    profile["branding"]["tab_colours"] = {"summary": "124481", "control": "71717A", "detail": "F4F4F5",
                                          "dashboard": "FF5D3B", "entity": "C4D9F0"}
    res = build_mod.build("demo", "2026-05", tmp_path / "b", api=FakeApi(), now=NOW)
    assert openpyxl.load_workbook(res.path)["Profit and loss"].sheet_properties.tabColor.rgb.endswith("124481")


def test_contents_counts_warnings_separately_from_passes(tmp_path, profile):
    res = build_mod.build("demo", "2026-05", tmp_path, api=FakeApi(), now=NOW)
    ws = openpyxl.load_workbook(res.path)["Contents"]
    summary = next(ws.cell(r, 1).value for r in range(1, 20) if str(ws.cell(r, 1).value).startswith("Controls:"))
    warns = sum(1 for c in res.controls if c.status == "warn")
    assert summary == f"Controls: {len(res.controls) - warns} passed, {warns} warning"


def test_cash_flow_year_to_date_check_compares_with_this_month_end(tmp_path, profile):
    res = build_mod.build("demo", "2026-05", tmp_path, api=FakeApi(), now=NOW)
    ws = openpyxl.load_workbook(res.path)["Cash flow"]
    chk = next(r for r in range(5, ws.max_row + 1) if str(ws.cell(r, 1).value).startswith("Check:"))
    assert "'Balance sheet'!C" in ws.cell(chk, 4).value      # YTD against the month end, not the prior month
    assert "'Balance sheet'!C" in ws.cell(chk, 3).value
    assert "'Balance sheet'!E" in ws.cell(chk, 5).value


def test_statements_freeze_at_g5_and_group_account_rows(tmp_path, profile):
    res = build_mod.build("demo", "2026-05", tmp_path, api=FakeApi(), now=NOW)
    wb = openpyxl.load_workbook(res.path)
    for name in ("Profit and loss", "Balance sheet", "Cash flow"):
        assert wb[name].freeze_panes == "G5"
    ws = wb["Profit and loss"]
    sales = next(r for r in range(5, ws.max_row + 1) if str(ws.cell(r, 1).value).strip() == "200 - Sales")
    assert ws.row_dimensions[sales].outlineLevel == 1 and ws.row_dimensions[sales].hidden
    assert [ws.cell(3, c).value for c in (3, 4, 5, 7)] == ["Month", "Year to date", "Prior year", "Month"]
    assert [ws.cell(4, c).value for c in (3, 4, 5, 7)] == ["May 2026", "May 2026", "Mar 2026", "Apr 2026"]
